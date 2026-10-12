import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.Models.InventarioFinal.Reporte import (
    DespachoSalidaResponse,
    ReporteSalidasConDespachosResponse,
)
from app.Reportes.ProductoTerminado import numero_salida_por_movimiento
from app.utils.dates import ZONA_BOLIVIA

COLOR_PRIMARIO = "1F4E79"
COLOR_TOTAL = "DCE6F1"
FORMATO_FECHA = "DD/MM/YYYY"
FORMATO_FECHA_HORA = "DD/MM/YYYY HH:MM"
FORMATO_ENTERO = "#,##0"

FUENTE_TITULO = Font(bold=True, size=14, color=COLOR_PRIMARIO)
FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
RELLENO_ENCABEZADO = PatternFill("solid", fgColor=COLOR_PRIMARIO)
RELLENO_TOTAL = PatternFill("solid", fgColor=COLOR_TOTAL)
LINEA = Side(style="thin", color="B0B0B0")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)


def _hora_local(fecha: datetime) -> datetime:
    return fecha.astimezone(ZONA_BOLIVIA).replace(tzinfo=None)


def _escribir_tabla(hoja, fila_inicio: int, columnas: list[tuple], filas: list, total: list | None = None) -> int:
    for indice, (titulo, _, _) in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_inicio, column=indice, value=titulo)
        celda.font = FUENTE_ENCABEZADO
        celda.fill = RELLENO_ENCABEZADO
        celda.border = BORDE
        celda.alignment = Alignment(vertical="center", wrap_text=True)

    fila_actual = fila_inicio
    for elemento in filas:
        fila_actual += 1
        for indice, (_, valor, formato) in enumerate(columnas, start=1):
            celda = hoja.cell(row=fila_actual, column=indice, value=valor(elemento))
            celda.border = BORDE
            if formato:
                celda.number_format = formato

    if total is not None:
        fila_actual += 1
        for indice, valor in enumerate(total, start=1):
            celda = hoja.cell(row=fila_actual, column=indice, value=valor)
            celda.font = Font(bold=True)
            celda.fill = RELLENO_TOTAL
            celda.border = BORDE
            formato = columnas[indice - 1][2]
            if formato and isinstance(valor, (int, float)):
                celda.number_format = formato

    return fila_actual


def _ajustar_anchos(hoja, anchos: list[int]) -> None:
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho


def construir_excel_salidas_producto_terminado(
    data: ReporteSalidasConDespachosResponse,
    generado_por: str | None = None,
) -> bytes:
    libro = Workbook()

    resumen = libro.active
    resumen.title = "Resumen"
    resumen["A1"] = "Salidas de producto terminado"
    resumen["A1"].font = FUENTE_TITULO
    resumen["A2"] = "Desde"
    resumen["B2"] = data.PeriodoInicio
    resumen["B2"].number_format = FORMATO_FECHA
    resumen["C2"] = "Hasta"
    resumen["D2"] = data.PeriodoFin
    resumen["D2"].number_format = FORMATO_FECHA
    resumen["A3"] = "Líneas"
    resumen["B3"] = "Todas" if data.TodasLasLineas else ", ".join(data.NombresLineas)
    resumen["A4"] = "Generado"
    resumen["B4"] = datetime.now(ZONA_BOLIVIA).replace(tzinfo=None)
    resumen["B4"].number_format = FORMATO_FECHA_HORA
    resumen["C4"] = generado_por or ""
    for referencia in ("A2", "C2", "A3", "A4"):
        resumen[referencia].font = Font(bold=True)

    fila = _escribir_tabla(
        resumen,
        6,
        [
            ("Línea", lambda l: l.NombreProducto, None),
            ("Unidad", lambda l: l.Unidad.capitalize(), None),
            ("Cantidad salida", lambda l: l.CantidadSalida, FORMATO_ENTERO),
            ("N° salidas", lambda l: l.NumeroSalidas, FORMATO_ENTERO),
        ],
        data.Lineas,
    )

    _escribir_tabla(
        resumen,
        fila + 2,
        [
            ("Código", lambda p: p.CodigoPresentacion, None),
            ("Producto", lambda p: p.NombrePresentacion, None),
            ("Unidad", lambda p: p.Unidad.capitalize(), None),
            ("Cantidad salida", lambda p: p.CantidadSalida, FORMATO_ENTERO),
            ("N° salidas", lambda p: p.NumeroSalidas, FORMATO_ENTERO),
            ("Stock actual", lambda p: p.CantidadActual, FORMATO_ENTERO),
        ],
        data.Presentaciones,
    )
    _ajustar_anchos(resumen, [16, 24, 14, 16, 12, 14])

    detalle = libro.create_sheet("Detalle")
    numero_salida = numero_salida_por_movimiento(data)
    columnas_detalle = [
        ("N° salida", lambda m: numero_salida[m.IdMovimientoProductoTerminado], None),
        ("Fecha", lambda m: _hora_local(m.FechaMovimiento), FORMATO_FECHA_HORA),
        ("Línea", lambda m: m.NombreProducto, None),
        ("Código", lambda m: m.CodigoPresentacion, None),
        ("Producto", lambda m: m.NombrePresentacion, None),
        ("Unidad", lambda m: m.Unidad.capitalize(), None),
        ("Cantidad", lambda m: m.Cantidad, FORMATO_ENTERO),
        ("Observación", lambda m: m.Observacion or "", None),
        ("Usuario", lambda m: f"{m.PrimerNombre} {m.ApellidoPaterno}", None),
        ("CI", lambda m: m.Ci, None),
        ("Rol", lambda m: m.NombreRol, None),
    ]
    ultima_fila = _escribir_tabla(detalle, 1, columnas_detalle, data.Movimientos)
    detalle.freeze_panes = "A2"
    detalle.auto_filter.ref = f"A1:{get_column_letter(len(columnas_detalle))}{ultima_fila}"
    _ajustar_anchos(detalle, [11, 17, 14, 12, 18, 11, 11, 40, 22, 12, 30])

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def construir_excel_despacho_salida(
    data: DespachoSalidaResponse,
    generado_por: str | None = None,
) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = f"Salida {data.IdDespacho}"

    hoja["A1"] = f"Salida de producto terminado N° {data.IdDespacho}"
    hoja["A1"].font = FUENTE_TITULO
    datos = [
        ("Fecha", _hora_local(data.FechaMovimiento), FORMATO_FECHA_HORA),
        ("Registrado por", f"{data.PrimerNombre} {data.ApellidoPaterno} - CI {data.Ci} - {data.NombreRol}", None),
        ("Observación", data.Observacion or "", None),
        ("Generado", datetime.now(ZONA_BOLIVIA).replace(tzinfo=None), FORMATO_FECHA_HORA),
        ("Generado por", generado_por or "", None),
    ]
    for fila, (etiqueta, valor, formato) in enumerate(datos, start=2):
        hoja.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        celda = hoja.cell(row=fila, column=2, value=valor)
        if formato:
            celda.number_format = formato

    _escribir_tabla(
        hoja,
        len(datos) + 3,
        [
            ("Código", lambda m: m.CodigoPresentacion, None),
            ("Producto", lambda m: m.NombrePresentacion, None),
            ("Línea", lambda m: m.NombreProducto, None),
            ("Unidad", lambda m: m.Unidad.capitalize(), None),
            ("Cantidad", lambda m: m.Cantidad, FORMATO_ENTERO),
        ],
        data.Items,
        total=["TOTAL", "", "", "", data.CantidadTotal],
    )
    _ajustar_anchos(hoja, [16, 40, 16, 12, 12])

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
